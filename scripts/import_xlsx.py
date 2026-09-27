"""One-time import of Patrick's workbook -> supabase/seed.sql.

    pip install openpyxl
    python scripts/import_xlsx.py "ArconER matches 2014 thru 2026.xlsx" > supabase/seed.sql

Warnings (unknown names, point mismatches) go to stderr. Review them, fix ALIASES, re-run.
"""
import re
import sys

import openpyxl

# Spreadsheet spelling -> canonical player name.
ALIASES = {
    "stelman": "Stelly",
    "hough": "Hougher",
    "matty": "Matty Ice",
    "matt": "Matty Ice",
    "feldman": "Feldy",
    "nordeen": "Nordo",
}

FIRST_YEAR = 2001


def canon(name):
    name = re.sub(r"\s+", " ", str(name)).strip()
    return ALIASES.get(name.lower(), name)


def s(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def q(v):
    if v is None:
        return "null"
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"


def warn(*a):
    print("WARN", *a, file=sys.stderr)


def split_names(cell):
    return [canon(n) for n in str(cell).split("/") if n.strip()]


def parse_simple(ws):
    """2014-2017: header row, roster columns, a 'Winner ...' row, then free-text notes."""
    roster, winner, notes, after_winner = [], None, [], False
    for a, _, c in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        a, c = s(a), s(c)
        if a and a.lower().startswith("winner"):
            w = a[6:].strip().lower()
            winner = w if w in ("arcon", "838") else None
            after_winner = True
        elif after_winner:
            if a:
                notes.append(a)
        else:
            roster += [(canon(a), "arcon")] * bool(a) + [(canon(c), "838")] * bool(c)
    return dict(roster=roster, winner=winner, notes="\n".join(notes), sessions=[], points=(None, None))


def parse_full(ws):
    """2018+: team sessions in columns A-E, singles in G-K, roster/points table under a '838 | Points | Arcon' header."""
    rows = [list(r[:11]) + [None] * (11 - len(r[:11])) for r in ws.iter_rows(values_only=True)]

    # Roster from the points table.
    team_of, roster = {}, []
    hdr = next(i for i, r in enumerate(rows) if r[6] == 838 and s(r[8]) == "Arcon")
    for r in rows[hdr + 1:]:
        if s(r[6]) == "Individual Total" or (r[6] is None and r[8] is None):
            break
        for cell, team in ((r[6], "838"), (r[8], "arcon")):
            if s(cell):
                for n in split_names(cell):
                    team_of[n] = team
                    roster.append((n, team))

    # Cup total + winner.
    arcon_pts = team2_pts = winner = None
    for r in rows:
        if s(r[6]) == "Team Arcon" and r[9] is not None:
            arcon_pts = r[9]
            if s(r[10]) == "WINNER":
                winner = "arcon"
        if s(r[6]) == "Team 838" and r[9] is not None:
            team2_pts = r[9]
            if s(r[10]) == "WINNER":
                winner = "838"

    def side_team(names, where):
        teams = {team_of.get(n) for n in names} - {None}
        if len(teams) != 1:
            warn(ws.title, where, names, "team unknown/mixed:", teams)
            return None
        return teams.pop()

    sessions = []

    # Team sessions: a header row ('4 Ball Match' / '2 man scramble'), then pair / vs / pair, closed by '<Label> Total'.
    cur = None
    for i, r in enumerate(rows):
        a = s(r[0])
        if a and a.lower() in ("4 ball match", "2 man scramble"):
            cur = dict(format="4-ball" if a.lower().startswith("4") else "2-man scramble", label=None, matches=[])
            sessions.append(cur)
        elif a and a.endswith("Total") and cur is not None and cur["label"] is None:
            cur["label"] = a.replace(" Total", "")
        elif a == "vs" and cur is not None:
            top, bot = rows[i - 1], rows[i + 1] if i + 1 < len(rows) else [None] * 11
            if not s(top[0]) or not s(bot[0]) or s(bot[0]).endswith("Total"):
                continue
            sides = []
            for row in (top, bot):
                names = split_names(row[0])
                sides.append(dict(names=names, team=side_team(names, f"{cur['format']} row {i}"),
                                  front=s(row[1]), back=s(row[2]), total=s(row[3]), pts=row[4] or 0))
            cur["matches"].append(sides)

    # Singles: names down column G from row 3 until the points header, consecutive pairs are opponents.
    singles = [r for r in rows[2:hdr] if s(r[6]) and s(r[6]) not in ("Players", "Individual")]
    single_session = dict(format="singles", label="Singles", matches=[])
    for a_row, b_row in zip(singles[0::2], singles[1::2]):
        if not any(a_row[7:11] + b_row[7:11]):
            continue  # listed but never played
        sides = []
        for row in (a_row, b_row):
            names = split_names(row[6])
            sides.append(dict(names=names, team=side_team(names, "singles"),
                              front=s(row[7]), back=s(row[8]), total=s(row[9]), pts=row[10] or 0))
        single_session["matches"].append(sides)
    if len(singles) % 2:
        warn(ws.title, "odd number of singles rows; last one dropped:", singles[-1][6])
    sessions.append(single_session)

    for n, sess in enumerate(sessions):
        sess["label"] = sess["label"] or sess["format"]
        sess["sort"] = n

    # Cross-check computed totals against the sheet.
    calc = {"arcon": 0, "838": 0}
    for sess in sessions:
        for sides in sess["matches"]:
            for side in sides:
                if side["team"]:
                    calc[side["team"]] += side["pts"]
    print(ws.title, "computed", calc, "sheet", arcon_pts, team2_pts, file=sys.stderr)
    if arcon_pts is not None and (calc["arcon"] != arcon_pts or calc["838"] != team2_pts):
        warn(ws.title, f"computed arcon {calc['arcon']} / 838 {calc['838']} vs sheet {arcon_pts} / {team2_pts}")

    return dict(roster=roster, winner=winner, notes=None, sessions=sessions, points=(arcon_pts, team2_pts))


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    years = {}
    for ws in wb:
        year = int(ws.title)
        years[year] = parse_full(ws) if ws.max_column > 3 else parse_simple(ws)

    players = sorted({n for y in years.values() for n, _ in y["roster"]}
                     | {n for y in years.values() for sess in y["sessions"] for m in sess["matches"] for side in m for n in side["names"]})
    last = max(years)

    # Compact SQL: multi-row VALUES, names resolved by join, ids derived from md5 of a short key.
    rows = lambda vals: ",\n".join("(" + ", ".join(q(v) for v in r) + ")" for r in vals)
    roster, sessions, matches, mplayers = [], [], [], []
    for year, y in sorted(years.items()):
        seen = set()
        played = [(n, side["team"]) for sess in y["sessions"] for m in sess["matches"] for side in m if side["team"] for n in side["names"]]
        for n, team in y["roster"] + played:
            if n not in seen:
                seen.add(n)
                roster.append((year, n, team))
        for sess in y["sessions"]:
            sid = f"s{year}-{sess['sort']}"
            sessions.append((sid, year, sess["label"], sess["format"], sess["sort"]))
            for k, sides in enumerate(sess["matches"]):
                by_team = {side["team"]: side for side in sides}
                if set(by_team) != {"arcon", "838"}:
                    warn(year, sess["label"], "skipping match, teams:", [(x["names"], x["team"]) for x in sides])
                    continue
                a_, b_ = by_team["arcon"], by_team["838"]
                mid = f"{sid}-{k}"
                matches.append((mid, sid, k, a_["front"], a_["back"], a_["total"], b_["front"], b_["back"], b_["total"], a_["pts"], b_["pts"]))
                mplayers += [(mid, n, side["team"]) for side in (a_, b_) for n in side["names"]]

    out = ["-- Generated by scripts/import_xlsx.py. Re-running replaces all imported data.",
           "delete from tournaments;", "delete from players;",
           "insert into players (name) values\n" + rows([(n,) for n in players]) + ";",
           "insert into tournaments (year, winner, arcon_points, team2_points, team2_name, notes) values\n" + rows(
               (yr, years.get(yr, {}).get("winner"), *(years[yr]["points"] if yr in years else (None, None)),
                "ER Systems" if yr <= 2011 else "838 Coatings", (years.get(yr, {}).get("notes") or None))
               for yr in range(FIRST_YEAR, last + 2)) + ";",
           "insert into roster (year, player_id, team) select v.y, p.id, v.t from (values\n" + rows(roster)
           + ") v(y, n, t) join players p on p.name = v.n;",
           "insert into sessions (id, year, label, format, sort) select md5(k)::uuid, y, l, f, o from (values\n" + rows(sessions)
           + ") v(k, y, l, f, o);",
           "insert into matches (id, session_id, sort, arcon_front, arcon_back, arcon_total, team2_front, team2_back, team2_total, arcon_pts, team2_pts)\n"
           "select md5(k)::uuid, md5(s)::uuid, o, af, ab, at, bf, bb, bt, ap, bp from (values\n" + rows(matches)
           + ") v(k, s, o, af, ab, at, bf, bb, bt, ap, bp);",
           "insert into match_players (match_id, player_id, team) select md5(v.k)::uuid, p.id, v.t from (values\n" + rows(mplayers)
           + ") v(k, n, t) join players p on p.name = v.n;"]
    print("\n".join(out))
    print(f"players: {len(players)}: {', '.join(players)}", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1])
